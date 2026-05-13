from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from pptx import Presentation
from pptx.util import Inches


def _chart_map(chart_paths: List[str]) -> Dict[str, str]:
    return {Path(path).name: path for path in chart_paths}


def export_board_pack(
    output_dir: str,
    company_name: str,
    period_label: str,
    kpis: Dict[str, Any],
    narrative: str,
    chart_paths: List[str],
) -> str:
    prs = Presentation()
    charts = _chart_map(chart_paths)

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"{company_name} CFO Board Pack"
    slide.placeholders[1].text = f"Monthly Performance | {period_label}"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive Summary"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    body.text = narrative

    p = body.add_paragraph()
    p.text = f"Revenue: ${kpis['revenue_latest']:,.0f} ({kpis['revenue_mom_pct'] * 100:+.1f}% MoM)"
    p.level = 0
    p = body.add_paragraph()
    p.text = f"Operating Margin: {kpis['operating_margin_latest'] * 100:.1f}%"
    p.level = 0
    p = body.add_paragraph()
    p.text = f"Cash Runway: {kpis['cash_runway_label']}"
    p.level = 0

    visuals = [
        ("Revenue Trend", charts.get("revenue_trend.png")),
        ("Operating Margin", charts.get("operating_margin.png")),
        ("Cash Runway", charts.get("cash_runway.png")),
        ("Budget vs Actual", charts.get("budget_vs_actual.png")),
    ]

    for title, chart_path in visuals:
        if not chart_path:
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_picture(chart_path, Inches(0.8), Inches(1.4), width=Inches(11.5))

    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    safe_period = period_label.replace(" ", "_")
    pptx_path = output / f"cfo_board_pack_{safe_period}.pptx"
    prs.save(str(pptx_path))
    return str(pptx_path)
